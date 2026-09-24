from datetime import datetime
from pathlib import Path
from core.models import MonitoringResult, MetricResult, Status
from reporting.production_reports import generate_system_excel, generate_system_pdf
import openpyxl

def test_system_excel_contains_operational_sheets(tmp_path):
    """
    The per-cycle workbook carries the operational sheets.

    This previously asserted a four-sheet redesign -- GUI_Evidence,
    Recovery, AI_Analysis -- and claimed in its own docstring to be
    "asserting what the code produces". It was not: generate_system_excel
    writes seven sheets under the original names, and the referenced
    TEST_FIXES.md does not exist in the repository. The test described a
    redesign that was never implemented, so it failed on every run and told
    the reader the opposite of the truth about the code.

    It now asserts the seven sheets actually written. If the four-sheet
    redesign is still wanted, that is a change to production_reports.py and
    this test moves with it.
    """
    result = MonitoringResult(system="TST", client="000", cycle_timestamp=datetime(2026,8,19,9,0))
    result.metrics.append(MetricResult(name="cpu", value=20, display_value="20%", status=Status.NORMAL))
    p = tmp_path / "TST_Monitoring.xlsx"
    generate_system_excel(result, [], str(p))
    wb = openpyxl.load_workbook(p)
    assert {"System Summary", "Metrics", "TCode Results", "Evidence Index",
            "Incidents", "Recovery History", "AI Analysis"} <= set(wb.sheetnames)
    # The metric itself must survive the round trip, not just its sheet.
    # Columns are located by header name rather than by position: the old
    # assertion hard-coded a three-column layout (Timestamp, Metric, Value)
    # that the writer has not produced for some time -- it writes eleven,
    # starting Timestamp, System, Client -- so it read the system name out
    # of column 2 and compared it to "cpu".
    metrics = wb["Metrics"]
    header = [c.value for c in metrics[1]]
    assert header[:5] == ["Timestamp", "System", "Client", "Metric", "Value"]
    assert metrics.cell(2, header.index("Metric") + 1).value == "cpu"
    assert metrics.cell(2, header.index("Value") + 1).value == 20
    assert metrics.cell(2, header.index("Status") + 1).value == "NORMAL"


def test_system_excel_honours_an_explicit_output_path(tmp_path):
    """output_path must override the derived reports/<date>/<system>/ location,
    including into a directory that does not exist yet -- a caller naming a
    path is not also promising to have created its parent."""
    result = MonitoringResult(system="TST", client="000", cycle_timestamp=datetime(2026,8,19,9,0))
    nested = tmp_path / "does" / "not" / "exist" / "TST.xlsx"
    written = generate_system_excel(result, [], str(nested))
    assert Path(written) == nested and nested.exists()

def test_system_pdf_contains_content(tmp_path):
    result = MonitoringResult(system="QAS", client="100", cycle_timestamp=datetime(2026,8,19,9,0))
    p = tmp_path / "QAS_Monitoring_Report.pdf"
    generate_system_pdf(result, [], str(p))
    assert p.exists()
    assert p.stat().st_size > 1000
