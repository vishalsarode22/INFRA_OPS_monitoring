"""
Round 18: screen readers checked against PS4's real layouts, and the
PDF/Excel reports checked for agreement.

The fake session below reproduces what SAP GUI scripting exposes for a
classic ABAP list: one label per cell at lbl[col,row], only the rows on the
current page, PageDown (VKey 82) / first page (VKey 80), a status bar, and
the Shift+F7 List Status popup. Column positions are taken from the PS4
screenshots of 22.09.2026.
"""
import time
from datetime import datetime

import pytest

import sap_gui.tcode_actions as ta
from core.models import MetricResult, MonitoringResult, Status


# --------------------------------------------------------------------------
# Fake SAP GUI
# --------------------------------------------------------------------------
class _Ctl:
    def __init__(self, id_="", text="", type_="GuiLabel", **attrs):
        self.Id, self.Text, self.Type = id_, text, type_
        self.__dict__.update(attrs)

    def press(self): pass
    Press = press

    def SetFocus(self): pass


class _Children:
    def __init__(self, items): self._items = list(items)
    @property
    def Count(self): return len(self._items)
    def __call__(self, i): return self._items[i]
    def __iter__(self): return iter(self._items)


class _Window:
    def __init__(self, session, index): self.s, self.i = session, index
    def sendVKey(self, key): self.s.vkey(self.i, key)
    def maximize(self): pass


class FakeSession:
    def __init__(self, pages=None, status_text="", records_passed=None, usr_children=None):
        self.pages = pages or [{}]
        self.page = 0
        self.status_text = status_text
        self.records_passed = records_passed
        self.popup = False
        self.keys = []
        self.usr_children = usr_children

    def vkey(self, window, key):
        self.keys.append((window, key))
        if window == 0 and key == 82:
            self.page = min(self.page + 1, len(self.pages) - 1)
        elif window == 0 and key == 80:
            self.page = 0
        elif window == 0 and key == 19 and self.records_passed is not None:
            self.popup = True
        elif window == 1 and key == 12:
            self.popup = False

    def findById(self, id_):
        if id_ in ("wnd[0]", "wnd[1]"):
            if id_ == "wnd[1]" and not self.popup:
                raise Exception("control not found")
            return _Window(self, int(id_[4]))
        if id_ == "wnd[0]/usr":
            if self.usr_children is not None:
                return _Ctl(id_, type_="GuiUserArea", Children=_Children(self.usr_children))
            cells = [_Ctl(f"/app/con[0]/ses[0]/wnd[0]/usr/lbl[{x},{y}]", text)
                     for (x, y), text in self.pages[self.page].items()]
            return _Ctl(id_, type_="GuiUserArea", Children=_Children(cells))
        if id_ == "wnd[0]/sbar":
            return _Ctl(id_, self.status_text, "GuiStatusbar", MessageType="S")
        if id_ == "wnd[1]/usr":
            if not self.popup:
                raise Exception("control not found")
            labels = [_Ctl("lbl1", "Records passed", CharTop=9, CharLeft=2),
                      _Ctl("lbl2", str(self.records_passed), CharTop=9, CharLeft=26)]
            return _Ctl(id_, Children=_Children(labels))
        if id_.startswith("wnd[0]/usr/cntl"):
            raise Exception("control not found")
        return _Ctl(id_)


@pytest.fixture(autouse=True)
def _fast(monkeypatch):
    monkeypatch.setattr(time, "sleep", lambda *_: None)
    monkeypatch.setattr(ta, "goto_tcode", lambda *a, **k: None)
    monkeypatch.setattr(ta, "wait_until_not_busy", lambda *a, **k: None)


def _capture_counter():
    shots = []

    def capture(suffix=""):
        shots.append(suffix)
        return None
    return capture, shots


# PS4 SM37 job overview: Spool and Job doc columns push Status to col 75.
_SM37_HEADER = {(4, 12): "JobName", (46, 12): "Spool", (52, 12): "Job doc",
                (62, 12): "Job CreatedB", (75, 12): "Status", (88, 12): "Start date",
                (99, 12): "Start Time", (110, 12): "Duration(sec.)",
                (126, 12): "Delay (sec.)", (138, 12): "Cli"}
_SM37_JOBS = [("/1DH/OBSERVE_LOGTAB", "11:26:55", "3.277")] + \
    [("/1LT/IUC_LOAD_MT_001_001", "12:10:01", "691")] * 5 + \
    [("/1LT/IUC_LOAD_MT_001_002", "12:14:30", "422")] * 3 + \
    [("/1LT/IUC_LOAD_MT_001_003", "12:16:58", "274"),
     ("/1LT/IUC_LOAD_MT_001_004", "12:15:56", "336")]


def _sm37_page(jobs, summary=False):
    page = dict(_SM37_HEADER)
    for i, (name, start, dur) in enumerate(jobs):
        y = 14 + i
        page.update({(4, y): name, (62, y): "SAP_WFRT", (75, y): "Active",
                     (88, y): "22.09.2026", (99, y): start, (118, y): dur,
                     (130, y): "0", (138, y): "500"})
    if summary:
        page[(4, 14 + len(jobs) + 1)] = "*Summary"
    return page


# --------------------------------------------------------------------------
# Screen readers
# --------------------------------------------------------------------------
def test_sm37_reads_status_by_header_across_pages_with_one_screenshot():
    session = FakeSession([_sm37_page(_SM37_JOBS[:8]), _sm37_page(_SM37_JOBS[8:], summary=True)],
                          records_passed=11)
    capture, shots = _capture_counter()
    data = ta.action_sm37(session, capture)
    assert data["active_job_count"] == 11
    assert data["records_passed"] == 11 and data["rows_read"] == 11
    assert data["longest_running_job"]["job_name"] == "/1DH/OBSERVE_LOGTAB"
    assert data["longest_running_job"]["duration_seconds"] == 3277
    assert shots == ["after_execution"], "the List Status popup must not be captured again"
    assert session.page == 0, "the list is returned to its first page"


def test_sm37_cancelled_reads_zero_from_the_status_bar():
    session = FakeSession(status_text="No job matches the selection criteria")
    capture, shots = _capture_counter()
    data = ta.action_sm37_cancelled(session, capture)
    assert data["cancelled_job_count"] == 0 and data["status"] == "OK"
    assert data["extraction_method"] == "sap_gui_status_bar"
    assert (0, 19) not in session.keys, "no List Status popup exists on the selection screen"
    assert len(shots) == 1


def test_sm37_cancelled_lists_names_when_jobs_exist():
    page = dict(_SM37_HEADER)
    page.update({(4, 14): "ZFI_DAILY_POST", (62, 14): "BATCHUSER", (75, 14): "Canceled",
                 (4, 15): "ZMM_STOCK_SYNC", (62, 15): "BATCHUSER", (75, 15): "Canceled",
                 (4, 17): "*Summary"})
    data = ta.action_sm37_cancelled(FakeSession([page], records_passed=2), _capture_counter()[0])
    assert data["cancelled_job_count"] == 2 and data["status"] == "WARNING"
    assert data["cancelled_job_names"] == ["ZFI_DAILY_POST", "ZMM_STOCK_SYNC"]


def test_smlg_reads_ps4_layout_and_grades_1566_ms_as_warning():
    from reporting.excel_template_writer import _status_from_smlg
    page = {(1, 3): "Application server", (22, 3): "State", (28, 3): "Resp.time(ms)",
            (42, 3): "Thrshd", (49, 3): "User", (54, 3): "Thrshd", (61, 3): "Time",
            (70, 3): "Quality", (78, 3): "Dialog steps",
            (1, 5): "vhrrnps4ai01_PS4_00", (35, 5): "1.149", (50, 5): "62",
            (61, 5): "12:18:33", (71, 5): "162", (86, 5): "39",
            (1, 6): "vhrrnps4ci_PS4_00", (35, 6): "1.566", (50, 6): "78",
            (61, 6): "12:20:32", (73, 6): "6", (86, 6): "81",
            (1, 8): "* Summary", (50, 8): "140",
            (1, 11): "Logon Group", (18, 11): "Current Instance",
            (1, 13): "PUBLIC", (18, 13): "vhrrnps4ai01_PS4_00"}
    data = ta.action_smlg(FakeSession([page]), lambda *a: "shot.png")
    assert [i["response_time_ms"] for i in data["instances"]] == [1149.0, 1566.0]
    assert data["response_time_ms"] == 1566.0, "the collector grades this into a metric"
    assert _status_from_smlg(data) == "WARNING"


def test_smlg_unreadable_is_not_measured():
    data = ta.action_smlg(FakeSession([{(1, 1): "CCMS: Load Distribution"}]), lambda *a: None)
    assert data["extraction_failed"] and data["instances"] == []


def _sp01_page(start, count, statuses=None):
    page = {(2, 3): "Spool no.", (12, 3): "Type", (17, 3): "User Name", (30, 3): "Date",
            (41, 3): "Time", (47, 3): "Status", (57, 3): "Pages", (64, 3): "Title"}
    for i in range(count):
        y = 5 + i
        page.update({(4, y): str(9400 + start + i), (17, y): "BATCHUSER",
                     (30, y): "22.09.2026", (47, y): (statuses or {}).get(start + i, "-"),
                     (60, y): "2", (64, y): "LIST1S Z54_CPI_STO_"})
    return page


@pytest.fixture
def _sp01_env(monkeypatch):
    monkeypatch.setattr(ta, "_wildcard_field_holding_logon_user", lambda s: True)
    monkeypatch.setattr(ta, "_labelled_fields", lambda s: [])


def test_sp01_counts_every_page_of_the_list(_sp01_env):
    session = FakeSession([_sp01_page(0, 30), _sp01_page(30, 15, {33: "Error"})])
    data = ta.action_sp01(session, lambda *a: None)
    assert data["spool_requests"] == 45
    assert data["spool_errors"] == 1
    assert data["spool_without_output_request"] == 44
    assert data["selection"]["created_by"] == "*"


def test_sp01_unreadable_list_is_unknown_not_zero(_sp01_env):
    data = ta.action_sp01(FakeSession([{(1, 1): "Output Controller"}]), lambda *a: None)
    assert data["spool_requests"] is None and data["extraction_failed"]


class _Grid(_Ctl):
    def __init__(self, rows):
        super().__init__("grid", type_="GuiShell", SubType="GridView")
        self.rows = rows
        self.RowCount, self.ColumnCount, self.VisibleRowCount = len(rows), 2, 10
        self.firstVisibleRow = 0
        self.Children = _Children([])

    def ColumnOrder(self, i): return ["TIME", "MSGTEXT"][i]
    def GetCellValue(self, r, c): return self.rows[r][0 if c == "TIME" else 1]


def test_sm21_finds_a_nested_grid_by_type():
    rows = [("10:06:40", "Delete ABAP session T72 [Warning/Session]")] * 7 + \
           [("10:00:06", "Stops work process 83")] * 25
    container = _Ctl("cont", type_="GuiContainer", Children=_Children([_Grid(rows)]))
    data = ta.action_sm21(FakeSession(usr_children=[container]), lambda *a: None)
    assert data["log_count"] == 32 and data["warning_count"] == 7
    assert data["extraction_method"] == "sap_gui_alv"


def test_sm21_unreadable_screen_is_a_failed_read():
    data = ta.action_sm21(FakeSession(usr_children=[]), lambda *a: None)
    assert data["extraction_failed"] and data["log_count"] is None


def test_db02_percentages_are_locale_safe():
    from utils.sap_numbers import parse_sap_decimal, parse_usage_ratio
    assert parse_sap_decimal("795,52") == 795.52
    assert parse_sap_decimal("1.149") == 1149
    assert parse_sap_decimal("1.234,56") == 1234.56
    assert parse_usage_ratio("795,52 GB /1,13 TB")["usage_percent"] == pytest.approx(68.75, abs=0.01)
    assert parse_usage_ratio("682,15 GB /912,34 GB")["usage_percent"] == pytest.approx(74.77, abs=0.01)


# --------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------
def _gui(tcode, shots=(), **data):
    data.setdefault("finished_at", "2026-09-22T12:21:10")
    data.setdefault("evidence_id", f"EV-PS4-{tcode}")
    return MetricResult(name=f"screenshot_{tcode}", value=None, display_value="captured (1)",
                        status=Status.UNKNOWN, source="sap_gui_collector", tcode=tcode,
                        screenshot_paths=list(shots), extra_data=data)


def test_old_false_zeros_now_read_unknown():
    from reporting.excel_template_writer import _status_for_metric, _build_actual_result
    sm21 = _gui("SM21", log_count=0, error_count=0, warning_count=0,
                extraction_method="sap_gui_labels", syslog_rows_count=32)
    sm37 = _gui("SM37", status="OK", active_job_count=0, records_passed=11,
                jobs=[{"job_name": "/1LT/X", "status": ""}])
    assert _status_for_metric(sm21) == "UNKNOWN"
    assert _status_for_metric(sm37) == "UNKNOWN"
    assert _build_actual_result(sm21).startswith("Not verified")
    smlg = _gui("SMLG", instance_count=0, instances=[], extraction_method="sap_gui_labels",
                ocr_lines=24)
    text = _build_actual_result(smlg)
    assert text.startswith("Not measured") and "Ocr Lines" not in text


def test_pdf_and_excel_use_the_same_status_words():
    from reporting.check_narratives import narrate_all
    from reporting.excel_template_writer import _status_for_metric
    gui = [_gui("SM58", trfc_status="CRITICAL", information={"entries_displayed": 3,
                "failed_entries": 3, "entries_in_execution": 0}),
           _gui("ST22", dump_count=3, dumps=[]), _gui("SCOT", mail_port=25, mail_port_raw="25"),
           _gui("SM37_CANCELLED", status="OK", cancelled_job_count=0)]
    for n, m in zip(narrate_all(gui), gui):
        assert n.status == _status_for_metric(m)
        assert n.status in {"OK", "WARNING", "CRITICAL", "UNKNOWN", "FAILED"}


def test_cross_checks_compare_like_with_like():
    from reporting.check_narratives import narrate_all, deterministic_analysis
    gui = [_gui("SM12", lock_count=35),
           _gui("SM58", trfc_status="CRITICAL", failed_entries=3,
                information={"entries_displayed": 3, "failed_entries": 3})]
    res = MonitoringResult(system="PS4", client="500")
    res.metrics = [
        MetricResult(name="sap.sm12.oldest_lock_minutes", value=29399, display_value="29399 min",
                     status=Status.CRITICAL),
        MetricResult(name="sap.sm58.stuck_entries", value=500, display_value="500 count",
                     status=Status.CRITICAL),
    ]
    res.overall_status = Status.CRITICAL
    da = deterministic_analysis(res, narrate_all(gui), gui)
    assert not any("oldest_lock" in c or "[SM12]" in c for c in da["conflicts"])
    assert any(c.startswith("[SM58]") and "500" in c for c in da["conflicts"])
    assert da["severity"] == "CRITICAL"


def test_pdf_shows_one_screenshot_per_check(tmp_path):
    from PIL import Image as PILImage
    from reporting.production_reports import _primary_screenshot, generate_system_pdf

    shade = iter(range(40, 250, 30))

    def png(name):
        # Distinct pixels: reportlab stores identical images only once.
        path = tmp_path / name
        PILImage.new("RGB", (1920, 1020), (next(shade), 200, 200)).save(path)
        return str(path)

    sm37 = [png("SM37_after_execution_1.png"), png("SM37_records_passed_2.png")]
    sp01 = [png("SP01_selection_1.png"), png("SP01_2.png")]
    assert _primary_screenshot(sm37) == sm37[0]
    assert _primary_screenshot(sp01) == sp01[1]

    gui = [_gui("SM37", sm37, status="OK", active_job_count=11, job_count=11),
           _gui("SP01", sp01, spool_requests=4, spool_errors=0),
           _gui("ST22", [png("ST22_today_1.png")], dump_count=0)]
    res = MonitoringResult(system="PS4", client="500", cycle_timestamp=datetime(2026, 9, 22, 12, 19))
    out = tmp_path / "r.pdf"
    generate_system_pdf(res, gui, str(out))
    raw = out.read_bytes()
    assert raw.count(b"/Subtype /Image") == 3


def test_excel_sheet_uses_capture_times_and_ok_in_guide(tmp_path, monkeypatch):
    import openpyxl
    import reporting.excel_template_writer as xw
    monkeypatch.setattr(xw, "_system_config", lambda name: {"sap_system_id": "PS4",
                                                            "environment": "Production"})
    gui = [_gui("AL08", user_logons="137", back_end_sessions="175",
                finished_at="2026-09-22T12:20:26")]
    res = MonitoringResult(system="PS4", client="500")
    out = xw.fill_metrobrands_template(
        "config/templates/InfraBeatOps_Simple_Monitoring_Sheet.xlsx", str(tmp_path / "s.xlsx"),
        gui, result=res)
    wb = openpyxl.load_workbook(out)
    ws = wb["Monitoring Sheet"]
    assert (ws["A4"].value, ws["B4"].value) == ("22.09.2026", "12:20:26")
    assert ws["D4"].value == "PS4" and ws["F4"].value == "Production"
    assert ws["I4"].value == "137 user sessions; 175 ABAP sessions"
    assert "OK" in [c.value for c in wb["Status Guide"]["A"]]
    assert "HEALTHY" not in [c.value for c in wb["Status Guide"]["A"]]
