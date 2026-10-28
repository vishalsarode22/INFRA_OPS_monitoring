"""
The new checks end to end, without SAP: recorded screens -> readers ->
grading -> monitoring sheet (rows follow the system's T-code set) -> PDF
(detail tables, every screen of the new checks) -> report gate.
"""
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import fake_sapgui as F  # noqa: E402
from sap_gui import tcode_actions_extra as X  # noqa: E402

S = F.screens()
F.install(X, {})
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SYSTEM = "TEST S4Q"


def _png(path):
    from PIL import Image
    Image.new("RGB", (480, 260), (230, 236, 245)).save(path)


class Capture:
    def __init__(self, tcode, folder):
        self.tcode, self.folder, self.paths = tcode, folder, []

    def __call__(self, suffix=""):
        name = f"{self.tcode}_{suffix}" if suffix else self.tcode
        p = os.path.join(self.folder, f"{name}_101500_{len(self.paths):06d}.png")
        _png(p)
        self.paths.append(p)
        return p


def _evidence(tcode, facts, shots):
    from core.models import MetricResult, Status
    data = {**facts, "evidence_id": f"EV-{tcode}", "attempts": 1, "system": SYSTEM, "client": "100",
            "started_at": "2026-10-13T10:15:00", "finished_at": "2026-10-13T10:15:09"}
    return MetricResult(name=f"screenshot_{tcode}", value=None, display_value=f"captured ({len(shots)})",
                        status=Status.UNKNOWN, source="sap_gui_collector", tcode=tcode, detail="",
                        screenshot_path=shots[0] if shots else None, screenshot_paths=list(shots),
                        extra_data=data)


class EndToEnd(unittest.TestCase):
    _built = False

    def setUp(self):
        # Per test (not setUpClass): some runners call TestCase.run() directly.
        from core import store, tcode_sets
        self._old = (store.DB_PATH, tcode_sets.PATH)
        if not EndToEnd._built:
            EndToEnd._build()
            EndToEnd._built = True
        store.DB_PATH = os.path.join(self.tmp, "t.sqlite")
        tcode_sets.PATH = os.path.join(self.tmp, "tcode_sets.yaml")

    def tearDown(self):
        from core import store, tcode_sets
        store.DB_PATH, tcode_sets.PATH = self._old

    @classmethod
    def _build(cls):
        from core import store, tcode_sets
        cls.tmp = tempfile.mkdtemp()
        if not os.environ.get("KEEP_TMP"):
            import atexit
            atexit.register(shutil.rmtree, cls.tmp, True)
        else:
            print("kept", cls.tmp)
        from sap_gui.tcode_actions import ACTIONS

        def run(tcode, action, session, settings=None):
            X.SETTINGS = settings or {}
            cap = Capture(tcode, cls.tmp)
            facts = ACTIONS[action](session, cap)
            X.SETTINGS = {}
            return _evidence(tcode, facts, cap.paths)

        # ST04 with large tables
        start = F.clone(S["st04_start"])
        tree = F.clone(S["st04_large"]["trees"][0])
        start["trees"] = [tree]
        s04 = F.Session(start)
        large = F.clone(S["st04_large"])
        large["trees"] = [tree]
        filled = F.clone(large)
        g = next(g for g in filled["grids"] if any(c["id"] == "TABLE_NAME" for c in g["columns"]))
        g["rows"] = [{"TABLE_NAME": f"TAB{i}", "SCHEMA_NAME": "SAPABAP1",
                      "MEMORY_SIZE_IN_TOTAL": f"{(20 - i) * 1024 ** 3}", "RECORD_COUNT": f"{i * 1000}"}
                     for i in range(12)]
        s04.on("open", lambda v: v[0] == "Large Tables", lambda s, v: s.show(F.clone(large)))
        s04.on("press", "Apply Selection", lambda s, v: s.show(F.clone(filled)))

        # STRUST
        st = F.clone(S["strust_start"])
        stree = F.clone(S["strust_pses"]["SSLSDFAULT1"]["trees"][0])
        st["trees"] = [stree]
        sst = F.Session(st)
        missing = F.clone(S["strust_pses"]["WSSEDFAULT"])

        def opened(s, v):
            sc = F.clone(S["strust_pses"].get(v[1], missing))
            sc["trees"] = [stree]
            s.show(sc, (sc.get("status") or "", "E" if sc.get("status") else ""))
        sst.on("open", None, opened)

        # ST05 all off
        st05 = F.clone(S["st05"])
        for n, v in (("STACK_OFF", True), ("PROGRESS_ON", True)):
            F.set_obj(st05, n, selected=v)

        cls.gui = [run("STRUST", "strust_certs", sst),
                   run("ST04", "st04_db", s04, {"top_n": 5}),
                   run("ST02", "st02_tune", F.Session(F.clone(S["st02"]))),
                   run("ST05", "st05_status", F.Session(st05))]

    def _set(self):
        from core import tcode_sets
        items = []
        on = ["SM12", "STRUST", "ST04", "ST02", "ST05"]
        for c in tcode_sets.catalogue():
            items.append({"tcode": c["tcode"], "on": c["tcode"] in on})
        # order: SM12 first, then the new ones as listed
        items.sort(key=lambda i: on.index(i["tcode"]) if i["tcode"] in on else 99)
        tcode_sets.save(SYSTEM, items)
        return on

    def test_sheet_pdf_and_gate(self):
        from core.models import MonitoringResult
        from core.tcode_metrics import normalize_tcode_results
        from core import report_gate, tcode_sets
        from reporting.excel_template_writer import fill_metrobrands_template
        from reporting.production_reports import generate_system_pdf
        import openpyxl

        order = self._set()
        self.assertEqual(tcode_sets.order_for(SYSTEM), order)
        self.assertEqual([t["tcode"] for t in tcode_sets.tasks_for(SYSTEM, ["ST04", "SM12", "SM50"])],
                         ["SM12", "ST04"])                      # a profile narrows, never adds
        result = MonitoringResult(system=SYSTEM, client="100")
        result.tcodes_this_run = order
        metrics = {m.name: m for m in normalize_tcode_results(self.gui, system=SYSTEM)}
        self.assertEqual(metrics["sap.st04.tenant_memory_pct"].status.value, "CRITICAL")
        self.assertEqual(metrics["sap.strust.min_days_left"].status.value, "NORMAL")
        self.assertEqual(metrics["sap.st05.traces_on"].value, 0.0)

        out = os.path.join(self.tmp, "sheet.xlsx")
        fill_metrobrands_template(os.path.join(ROOT, "config/templates/InfraBeatOps_Simple_Monitoring_Sheet.xlsx"),
                                  out, self.gui, result=result)
        ws = openpyxl.load_workbook(out)["Monitoring Sheet"]
        rows = [[ws.cell(r, c).value for c in range(1, 12)] for r in range(4, 24)]
        self.assertEqual([r[6] for r in rows[:5]], order)
        self.assertTrue(all(all(v is None for v in r) for r in rows[5:]))       # template rows emptied
        self.assertEqual(ws["A2"].value, "5 checks")
        by = {r[6]: r for r in rows[:5]}
        self.assertEqual(by["ST04"][9], "CRITICAL")
        self.assertIn("HANA memory 89.7%", by["ST04"][8])
        self.assertIn("tenant memory", by["ST04"][10].lower())
        self.assertEqual(by["STRUST"][9], "OK")
        self.assertIn("8 own certificates", by["STRUST"][8])
        self.assertEqual(by["ST05"][9], "OK")
        self.assertEqual(by["ST05"][8], "All traces off; progress display on")
        self.assertIn(by["ST02"][9], ("WARNING", "CRITICAL"))          # CUA 7.1 million swaps since 22.08.
        self.assertEqual(by["ST04"][7], "Database Memory / Volumes / Large Tables")

        gaps, err = report_gate.sheet_gaps(out)
        self.assertEqual(err, "")
        self.assertEqual(gaps, ["SM12"])                                # no SM12 capture in this test
        self.assertEqual(report_gate._expected_tcodes({"name": SYSTEM, "_tcodes_this_run": order}), order)
        self.assertEqual(report_gate._expected_tcodes({"name": SYSTEM}), order)

        pdf = os.path.join(self.tmp, "r.pdf")
        generate_system_pdf(result, self.gui, output_path=pdf)
        text = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
        self.assertIn("Figures read from the screens", text)
        self.assertIn("Top 1 column table by memory", text)
        self.assertIn("SSL server Standard / System-wide / host1_S4Q_10", text)
        self.assertIn("large tables by memory", text)                   # every ST04 screen in the appendix
        self.assertEqual(text.count("STRUST — Certificate Expiry (pse"), 8)
        self.assertNotIn("SM13", text)                                  # unticked: not in the report

    def test_system_without_sap_gui(self):
        from core import config_loader, tcode_sets
        from core.models import MonitoringResult
        from reporting.excel_template_writer import fill_metrobrands_template
        import openpyxl
        old = config_loader.get_systems
        config_loader.get_systems = lambda: [{"name": "RFC ONLY", "has_gui_access": False}]
        try:
            items = {i["tcode"]: i["on"] for i in tcode_sets.for_system("RFC ONLY")}
            self.assertFalse(any(items[t] for t in tcode_sets.GUI_ONLY))      # screen-only checks start unticked
            self.assertTrue(items["SM12"])
            result = MonitoringResult(system="RFC ONLY", client="100")
            result.tcodes_this_run = ["SM12", "ST04"]                           # ticked by hand anyway
            out = os.path.join(self.tmp, "nogui.xlsx")
            fill_metrobrands_template(os.path.join(ROOT, "config/templates/InfraBeatOps_Simple_Monitoring_Sheet.xlsx"),
                                      out, [], result=result)
            ws = openpyxl.load_workbook(out)["Monitoring Sheet"]
            self.assertEqual((ws.cell(5, 7).value, ws.cell(5, 10).value), ("ST04", "NOT APPLICABLE"))
        finally:
            config_loader.get_systems = old


if __name__ == "__main__":
    unittest.main()
