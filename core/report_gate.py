"""
Send the monitoring sheet only when the monitoring really happened.

On 12 Oct 2026 a "Monitor this system" run for CEAT S4P could not log on to
SAP GUI (wrong credentials), took no screenshots, left 8 of 19 sheet rows
"Not captured / UNKNOWN" -- and the PDF and Excel still went to the client.

COMPLETE means all of:
    1. the SAP GUI logon worked (for a system with has_gui_access);
    2. every T-code of this run (the system's set on the Profile page) has a captured
       screenshot that exists on disk;
    3. every row of the Excel monitoring sheet has a status (OK, WARNING,
       CRITICAL, NOT AUTHORIZED or NOT APPLICABLE) -- no UNKNOWN, no empty
       status, no "Not captured".
A system can list T-codes that are allowed to be missing:
    report_gate: {allow_missing: [DB12]}        (config/systems.yaml)

NOT COMPLETE: nothing goes to the client. Basis gets a short internal mail
(no attachments) saying what is missing, and:
    * SAP GUI logon failed -> NO automatic retry. It is most often a wrong
      password, and retrying could lock the SAP user. Fix it, then press
      "Monitor this system".
    * anything else (screens not captured, sheet not filled, another SAP GUI
      session open on the VM) -> retried every IBO_REPORT_RETRY_MINUTES
      (15) up to IBO_REPORT_RETRY_MAX times (default 2) by the dashboard scheduler.
      After the last try, a run that logged on and misses only a few checks
      (IBO_REPORT_PARTIAL_MAX, default 2) is SENT, the missing checks named in
      the mail (28.10.2026: a PS4 sheet was lost over one empty ST02 row).
      After the last try Basis gets "not sent".

COMPLETE: the report goes to the client as before, and Basis gets an
internal copy notice with an "Acknowledge: report checked" button (a
signed, single-use link, core/action_links.py; it works on the company
network). Pressing it records who checked the report on the Timeline.

IBO_REPORT_GATE=0 switches this off (old behaviour: always send).
"""
from __future__ import annotations

import os
import re
from datetime import datetime, timedelta

from utils.logger import get_logger

log = get_logger(__name__)

GOOD_STATUSES = {"OK", "WARNING", "CRITICAL", "NOT AUTHORIZED", "NOT APPLICABLE"}


def enabled() -> bool:
    return os.getenv("IBO_REPORT_GATE", "1").strip().lower() not in ("0", "false", "no", "off")


def _int_env(name: str, default: int) -> int:
    try:
        return max(1, int(os.getenv(name, str(default)) or default))
    except ValueError:
        return default


def retry_minutes() -> int:
    return _int_env("IBO_REPORT_RETRY_MINUTES", 15)


def retry_max() -> int:
    """Extra automatic sweeps after a held report (default 2, 15 min apart).
    It was 0 until 28.10.2026: the 12:00 PS4 sheet was then held over one
    empty ST02 row and never tried again -- nobody got a report."""
    try:
        return max(0, int(os.getenv("IBO_REPORT_RETRY_MAX", "2") or 0))
    except ValueError:
        return 2


def partial_max() -> int:
    """After the last try: send anyway when at most this many checks are missing (0 = never)."""
    try:
        return max(0, int(os.getenv("IBO_REPORT_PARTIAL_MAX", "2") or 0))
    except ValueError:
        return 2


def partial_ok(gate: dict) -> bool:
    """
    The last try is still incomplete, but SAP GUI logged on, the sheet was
    built, and only a few checks are missing: the report goes out with those
    checks named, instead of nobody getting a report at all.
    """
    if not gate or gate.get("complete") or not gate.get("retry") or gate.get("no_retry"):
        return False
    if gate.get("sheet_error") or gate.get("gui_failed"):
        return False
    missing = list(gate.get("missing") or [])
    return 0 < len(missing) <= partial_max()


def partial_note(gate: dict) -> str:
    """The plain-language note at the top of a report sent with missing checks."""
    from html import escape as e
    det = gate.get("details") or {}
    items = "".join(f"<li><b>{e(t)}</b>: {e(det.get(t) or 'not captured in this run')}</li>"
                    for t in gate.get("missing") or [])
    return ("<div style='border:1px solid #f59e0b;background:#fffbeb;border-radius:8px;padding:10px 14px;"
            "margin:0 0 14px;font:13px Segoe UI,Arial,sans-serif;color:#78350f'>"
            f"<b>Note:</b> {len(gate.get('missing') or [])} check(s) could not be captured after "
            f"{retry_max() + 1} tries and are marked in the sheet. Everything else was checked as usual."
            f"<ul style='margin:6px 0 0 18px;padding:0'>{items}</ul></div>")


# ---------------------------------------------------------------- the check

def _expected_tcodes(system_config: dict | None = None) -> list[str]:
    """This run's T-codes (main._tasks_this_run), else the system's set from the Profile page."""
    cfg = system_config or {}
    if cfg.get("_tcodes_this_run"):
        return [str(t).strip().upper() for t in cfg["_tcodes_this_run"]]
    try:
        from core import tcode_sets
        return tcode_sets.order_for(str(cfg.get("name") or ""))
    except Exception as exc:  # noqa: BLE001
        log.warning(f"report gate: T-code set unreadable ({exc}); every monitoring task expected")
    try:
        from core.config_loader import get_monitoring_tasks
        return [str(t.get("tcode") or "").strip().upper() for t in get_monitoring_tasks() if t.get("tcode")]
    except Exception as exc:  # noqa: BLE001
        log.warning(f"report gate: monitoring tasks unreadable: {exc}")
        return []


def sheet_gaps(sheet_path: str, texts: dict | None = None) -> tuple[list[str], str]:
    """T-codes whose sheet row has no usable status, and an error text when the sheet cannot be read."""
    if not sheet_path or not os.path.exists(sheet_path):
        return [], "the Excel monitoring sheet was not created"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(sheet_path, read_only=True, data_only=True)
        ws = wb["Monitoring Sheet"] if "Monitoring Sheet" in wb.sheetnames else wb.worksheets[0]
        header, gaps = None, []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if header is None:
                if "T-Code" in cells and "Status" in cells:
                    header = {name: i for i, name in enumerate(cells)}
                continue
            tcode = cells[header["T-Code"]] if header["T-Code"] < len(cells) else ""
            if not tcode:
                continue
            status = (cells[header["Status"]] if header["Status"] < len(cells) else "").upper()
            actual = cells[header["Actual Result"]] if "Actual Result" in header and \
                header["Actual Result"] < len(cells) else ""
            if status not in GOOD_STATUSES or actual.lower().startswith("not captured"):
                gaps.append(tcode.upper())
                if texts is not None:
                    texts[tcode.upper()] = actual
        wb.close()
        if header is None:
            return [], "the Excel monitoring sheet has no T-Code / Status columns"
        return gaps, ""
    except Exception as exc:  # noqa: BLE001
        return [], f"the Excel monitoring sheet could not be read ({type(exc).__name__})"


# ---------------------------------------------------------------- why a check is missing

def _metrics_by_tcode(gui_results) -> dict:
    """The screenshot_<TCODE> result per T-code (the one its sheet row is built from)."""
    out = {}
    for r in gui_results or []:
        name = str(getattr(r, "name", "") or "")
        t = str(getattr(r, "tcode", "") or "").strip().upper()
        if not t:
            continue
        if name.startswith("screenshot_") or t not in out:
            out[t] = r
    return out


def _short(text, n=140) -> str:
    text = " ".join(str(text or "").split())
    text = re.sub(r"(?i)(passw\w*|pwd|bcode)\s*[=:]\s*\S+", r"\1=***", text)
    return text if len(text) <= n else text[: n - 3] + "..."


_STEP = {
    "goto": "the transaction did not open",
    "selection": "the selection screen could not be filled",
    "execute": "the selection was not executed (F8 failed)",
    "rejected": "SAP rejected the selection",
    "popup": "a popup blocked the screen",
    "busy": "SAP was still loading the list when it was read",
    "list": "the result list could not be read",
    "partial": "the list was only partly read",
    "exception": "the SAP GUI step failed",
    "stopped": "the run was stopped before this check",
}
_AUTH_RX = re.compile(r"not authori[sz]ed|no authori[sz]ation|keine berechtigung", re.I)


def reason_for(metric, sheet_text: str = "") -> str:
    """
    Plain-language reason why a check's sheet row is not filled, with the
    number of tries and SAP's last status-bar message when known.
    """
    if metric is None:
        if sheet_text and not sheet_text.lower().startswith(("see attached", "not captured")):
            return _short(sheet_text)
        return "no result came back from SAP GUI for this check (it did not run, or the run stopped before it)"
    d = getattr(metric, "extra_data", None) or {}
    sbar = _short(d.get("status_message") or d.get("last_status_bar") or "", 120)
    popup = d.get("popup") or ""
    if isinstance(d.get("popups"), list) and d["popups"]:
        popup = popup or d["popups"][0]
    popup = _short(popup, 100)
    step = str(d.get("failed_step") or "")
    err = _short(d.get("error") or d.get("extraction_error") or "")

    if _AUTH_RX.search(sbar) or _AUTH_RX.search(popup) or str(getattr(metric, "display_value", "")).upper() == "NOT AUTHORIZED":
        why = "the SAP user is not authorised for this screen"
    elif str(getattr(metric, "display_value", "")) == "failed":
        why = _STEP["exception"] + (f": {err}" if err else
                                    (f" ({_short(getattr(metric, 'detail', ''), 100)})"
                                     if getattr(metric, "detail", "") else ""))
    elif step in _STEP:
        why = _STEP[step] + (f": {err}" if err and step != "exception" else "")
    elif err:
        why = err
    elif sheet_text and sheet_text.lower().startswith("not measured"):
        why = _short(sheet_text[len("not measured"):].lstrip(" :"), 160) or "the values could not be read"
    elif d.get("extraction_failed"):
        why = "the screen opened but its values could not be read"
    else:
        why = "the screen was captured but no figure could be read from it"
    if popup and "popup" not in why:
        why += f"; popup: '{popup}'"
    extras = []
    try:
        attempts = int(d.get("attempts") or 0)
    except (TypeError, ValueError):
        attempts = 0
    if attempts:
        extras.append(f"{attempts} attempt{'s' if attempts != 1 else ''}")
    if sbar and sbar not in why:
        extras.append(f"SAP said: '{sbar}'")
    return why + (f" ({'; '.join(extras)})" if extras else "")


def check(system_config: dict, gui_available: bool, gui_skip_reason: str, gui_results: list,
          sheet_path: str) -> dict:
    """
    {"complete": bool, "reasons": [str], "retry": bool, "missing": [tcode]}.
    `retry` is False when another automatic try would not help or could harm
    (SAP GUI logon refused: possibly a wrong password).
    """
    reasons, missing = [], set()
    allow = {str(t).strip().upper() for t in
             ((system_config.get("report_gate") or {}).get("allow_missing") or [])}
    expected = [t for t in _expected_tcodes(system_config) if t not in allow]
    has_gui = system_config.get("has_gui_access", True)
    retry = True
    no_retry = ""
    details: dict = {}
    by_t = _metrics_by_tcode(gui_results)

    gui_failed = False
    if has_gui and not gui_available:
        gui_failed = True
        reasons.append(f"SAP GUI was not used: {gui_skip_reason or 'no SAP GUI session'}")
        missing.update(expected)
        why = (gui_skip_reason or "").lower()
        if "check the user, password" in why:
            retry, no_retry = False, "logon"           # do not hammer the SAP user
        elif "not available on this machine" in why:
            retry, no_retry = False, "no_gui"          # no SAP GUI here (Ubuntu)
    elif has_gui:
        captured = {}
        for r in gui_results or []:
            name = str(getattr(r, "name", "") or "")
            if not name.startswith("screenshot_"):
                continue
            paths = list(getattr(r, "screenshot_paths", None) or [])
            if getattr(r, "screenshot_path", None):
                paths.append(r.screenshot_path)
            ok = str(getattr(r, "display_value", "") or "").startswith("captured") and \
                any(p and os.path.exists(p) for p in paths)
            captured[name[len("screenshot_"):].upper()] = ok
        not_shot = [t for t in expected if not captured.get(t)]
        if not_shot:
            reasons.append(f"screenshots not captured for {len(not_shot)} T-code(s): {', '.join(not_shot)}")
            missing.update(not_shot)
            for t in not_shot:
                details.setdefault(t, reason_for(by_t.get(t)))

    gap_text: dict = {}
    gaps, sheet_error = sheet_gaps(sheet_path, gap_text)
    gaps = [g for g in gaps if g not in allow]
    if sheet_error:
        reasons.append(sheet_error)
    elif gaps:
        reasons.append(f"the Excel sheet is not filled for {len(gaps)} check(s): {', '.join(gaps)}")
        missing.update(gaps)
        for g in gaps:
            details[g] = reason_for(by_t.get(g), gap_text.get(g, ""))

    return {"complete": not reasons, "reasons": reasons, "retry": retry, "no_retry": no_retry,
            "missing": sorted(missing), "details": details, "sheet_error": bool(sheet_error),
            "gui_failed": gui_failed}


def _detail_lines(gate: dict) -> list[str]:
    return [f"{t} — not filled: {r}" for t, r in (gate.get("details") or {}).items()]


# ---------------------------------------------------------------- retries
#
# report_retries holds at most one chain per system:
#     attempt     the try that was last held back (1 = the first run)
#     due_at      when the next try should start
#     started_at  set when the dashboard scheduler starts that try; cleared
#                 when the try reaches the report step (hold / sent)
# A try that was started but never reached the report step (crash, hang)
# still counts: it is found the next time the scheduler looks, and the
# chain goes on or ends with a "not sent" mail -- never silently.

ALIVE_KEY = "report_retry_scheduler_alive"

# Set by the dashboard around a run a person started ("Monitor this system",
# "Run all"). Only one sweep runs per machine at a time (core/run_lock.py),
# so a module flag is enough; the sweep may run on a watchdog thread.
_manual_run = {"on": False}


def set_manual_run(on: bool) -> None:
    _manual_run["on"] = bool(on)


def manual_run() -> bool:
    return _manual_run["on"]


def _row(system: str) -> dict | None:
    from core import ops_db
    return ops_db.one("SELECT * FROM report_retries WHERE system=?", (system,))


def _fmt(d: datetime) -> str:
    return d.strftime("%Y-%m-%d %H:%M:%S")


def _parse(s) -> datetime | None:
    try:
        return datetime.strptime(str(s)[:19], "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return None


def scheduler_alive(now: datetime | None = None) -> bool:
    """
    True when automatic monitoring was on recently. The dashboard scheduler
    marks it on every pass while it is enabled; no pass happens while a sweep
    runs, so the window covers the longest sweep (IBO_REPORT_SCHEDULER_WINDOW_MIN,
    default 180).
    """
    from core import ops_db
    try:
        row = ops_db.one("SELECT value FROM meta WHERE key=?", (ALIVE_KEY,))
    except Exception:  # noqa: BLE001
        return False
    seen = _parse(row["value"]) if row else None
    window = _int_env("IBO_REPORT_SCHEDULER_WINDOW_MIN", 180)
    return bool(seen and (now or datetime.now()) - seen < timedelta(minutes=window))


def mark_scheduler_alive(now: datetime | None = None, paused: bool = False) -> None:
    """Called on every scheduler pass: the time while monitoring is on, "paused" while it is off."""
    from core import ops_db
    try:
        ops_db.execute("INSERT INTO meta(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET "
                       "value=excluded.value", (ALIVE_KEY, "paused" if paused else _fmt(now or datetime.now())))
    except Exception as exc:  # noqa: BLE001
        log.debug(f"report gate: scheduler heartbeat not stored: {exc}")


def next_attempt(system: str, now: datetime | None = None, manual: bool = False) -> int:
    """
    Which try this run is. A waiting chain less than 3 hours old is continued
    by any run (a scheduler retry, or a profile sweep that came first), so it
    always ends after 1 + IBO_REPORT_RETRY_MAX tries. A manual run
    (manual=True) starts a new chain: a person decided to try again.
    """
    now = now or datetime.now()
    row = _row(system)
    if not row or manual:
        return 1
    last = _parse(row.get("started_at")) or _parse(row.get("due_at"))
    if not last or now - last > timedelta(hours=3):
        return 1
    return int(row["attempt"]) + 1


def plan_retry(system: str, attempt: int, reason: str, now: datetime | None = None,
               profile_id: str | None = None) -> datetime | None:
    """Store the next try (returns its time), or None when the tries are used up.
    `profile_id`: the profile whose run was held (kept for every try of the chain)."""
    from core import ops_db
    now = now or datetime.now()
    if attempt >= retry_max() + 1:
        clear(system)
        return None
    due = now + timedelta(minutes=retry_minutes())
    if profile_id and str(profile_id).startswith("Retry"):
        profile_id = None
    ops_db.execute("INSERT INTO report_retries(system, attempt, due_at, reason, first_at, started_at, profile_id) "
                   "VALUES (?,?,?,?,?,NULL,?) ON CONFLICT(system) DO UPDATE SET attempt=excluded.attempt, "
                   "due_at=excluded.due_at, reason=excluded.reason, started_at=NULL, "
                   "profile_id=COALESCE(excluded.profile_id, report_retries.profile_id)",
                   (system, attempt, _fmt(due), reason[:900], _fmt(now), profile_id))
    return due


def retry_profile(system: str) -> str | None:
    """The profile whose held run this retry chain belongs to (None: a Scheduler-page/manual run)."""
    r = _row(system)
    return (r or {}).get("profile_id") or None


def unmark_started(system: str, minutes: int = 5) -> None:
    """The retry could not start (another run held SAP GUI): it does not count as a try."""
    from core import ops_db
    ops_db.execute("UPDATE report_retries SET started_at=NULL, due_at=? WHERE system=?",
                   (_fmt(datetime.now() + timedelta(minutes=minutes)), system))


def clear(system: str) -> None:
    from core import ops_db
    ops_db.execute("DELETE FROM report_retries WHERE system=?", (system,))


def mark_started(system: str, now: datetime | None = None) -> None:
    """The scheduler starts a retry. Its due time moves on, so a crash cannot restart it every pass."""
    from core import ops_db
    now = now or datetime.now()
    ops_db.execute("UPDATE report_retries SET started_at=?, due_at=? WHERE system=?",
                   (_fmt(now), _fmt(now + timedelta(minutes=retry_minutes())), system))


def postpone(system: str, minutes: int = 5) -> None:
    """A due retry that could not start (maintenance, parked): try again later."""
    from core import ops_db
    ops_db.execute("UPDATE report_retries SET due_at=? WHERE system=?",
                   (_fmt(datetime.now() + timedelta(minutes=minutes)), system))


def due_retries(now: datetime | None = None, system_configs: dict | None = None) -> list[str]:
    """
    Systems whose next try is due. Call only while no sweep is running: a row
    that was started but is still marked started then never reached the
    report step (crashed or hung), so that try counts as failed here.
    """
    from core import ops_db
    now = now or datetime.now()
    out = []
    for r in ops_db.rows("SELECT * FROM report_retries WHERE due_at<=? ORDER BY due_at", (_fmt(now),)):
        cfg = (system_configs or {}).get(r["system"]) or {"name": r["system"]}
        due_at = _parse(r["due_at"])
        if not r.get("started_at") and due_at and now - due_at > timedelta(hours=3):
            # Waited far past its time (automatic monitoring was off): a sweep
            # now would surprise people. End the chain and say so.
            clear(r["system"])
            prev = str(r.get("reason") or "").strip()
            _tell_basis(cfg, {"reasons": ["the retry could not run in time (automatic monitoring was off)"]
                              + ([f"last known cause: {prev}"] if prev else []),
                              "retry": True}, None, retry_max() + 1, now)
            continue
        if r.get("started_at"):
            attempt = int(r["attempt"]) + 1
            reason = "the retry did not finish (the run stopped before the report step)"
            if attempt >= retry_max() + 1:
                clear(r["system"])
                prev = str(r.get("reason") or "").strip()
                _tell_basis(cfg, {"reasons": [reason] + ([f"last known cause: {prev}"] if prev else []),
                                  "retry": True}, None, attempt, now)
                continue
            plan_retry(r["system"], attempt, reason, now - timedelta(minutes=retry_minutes()))
        out.append(r["system"])
    return out


# ---------------------------------------------------------------- mails to Basis

def _basis(system_config: dict) -> list[str]:
    """Basis addresses; with IBO_INTERNAL_DOMAINS / IBO_INTERNAL_ADDRESSES set, only those."""
    try:
        from core import recipients
        addrs = list(recipients.basis(system_config) or [])
    except Exception:  # noqa: BLE001
        addrs = []
    from core import action_links
    if action_links._domains() or action_links._addresses():
        addrs = [a for a in addrs if action_links.is_internal(a)]
    return addrs


def internal_smtp(system_config: dict, smtp_config: dict) -> dict:
    """A copy of smtp_config that sends to Basis only (failure notices while a report is held)."""
    out = dict(smtp_config or {})
    out["to_emails"] = _basis(system_config)
    out["cc_emails"] = []
    out["manager_emails"] = []
    return out


def _send_internal(system_config: dict, subject: str, html: str, kind: str, key: str,
                   severity: str) -> bool:
    from core.config_loader import get_smtp_config
    from notifications.mailer import send_html
    to = _basis(system_config)
    if not to:
        log.warning(f"{system_config.get('name')}: report gate: no Basis address to tell")
        return False
    return bool(send_html(subject=subject, html=html, to=to,
                          smtp_config=get_smtp_config(system_config, require_recipients=False),
                          system=system_config.get("name", ""), kind=kind, key=key, severity=severity))


def hold(system_config: dict, gate: dict, attempt: int, now: datetime | None = None) -> dict:
    """The report is NOT sent to the client. Plan a retry (when one can run) and tell Basis."""
    name = system_config.get("name", "")
    now = now or datetime.now()
    can_retry = gate["retry"] and scheduler_alive(now)
    due = plan_retry(name, attempt, "; ".join(gate["reasons"] + _detail_lines(gate)), now) if can_retry else None
    if not can_retry:
        clear(name)
    _tell_basis(system_config, gate, due, attempt, now)
    return {"held": True, "retry_at": _fmt(due) if due else None}


def _tell_basis(system_config: dict, gate: dict, due, attempt: int, now: datetime) -> None:
    from html import escape as e
    name = system_config.get("name", "")
    if due:
        what = (f"It will be tried again automatically at {due:%H:%M} "
                f"(try {attempt + 1} of {retry_max() + 1}).")
        head = "Monitoring sheet held back -- retrying"
    elif gate.get("no_retry") == "no_gui":
        what = ("Not retried: SAP GUI is not available on this machine, so screenshots cannot be taken "
                "here. Run the monitoring from the Windows monitoring VM.")
        head = "Monitoring sheet NOT sent"
    elif not gate.get("retry"):
        what = ("Not retried automatically: the SAP GUI logon was refused, often a wrong password, and "
                "repeated tries could lock the SAP user. Check the user and password, then press "
                "'Monitor this system'.")
        head = "Monitoring sheet NOT sent"
    elif gate.get("retry") and attempt < retry_max() + 1:
        what = ("Not retried automatically: the dashboard scheduler is not running (automatic "
                "monitoring paused, or the run was started with python main.py). Press "
                "'Monitor this system' when the cause is fixed.")
        head = "Monitoring sheet NOT sent"
    elif retry_max() == 0:
        what = ("Automatic retries are off, so it is not run again before its next scheduled slot. "
                "Fix the cause and press 'Monitor this system' to send it now.")
        head = "Monitoring sheet NOT sent"
    else:
        what = (f"All {retry_max() + 1} tries failed. Fix the cause and press 'Monitor this system'.")
        head = "Monitoring sheet NOT sent"
    html = ("<html><body style='font-family:Segoe UI,Arial,sans-serif;color:#0f172a'>"
            f"<h2 style='color:#b45309;margin:0 0 10px'>{e(head)}: {e(name)}</h2>"
            "<p>The monitoring run was not complete, so the PDF and Excel were <b>not</b> sent to the "
            "client.</p><ul>" + "".join(f"<li>{e(r)}</li>" for r in gate["reasons"]) + "</ul>"
            + (("<p style='margin:6px 0 2px'><b>Why:</b></p><ul>"
                + "".join(f"<li><b>{e(t)}</b> &mdash; {e(r)}</li>" for t, r in (gate.get("details") or {}).items())
                + "</ul>") if gate.get("details") else "")
            + f"<p>{e(what)}</p>"
            f"<p style='color:#64748b'>Run of {now:%d.%m.%Y %H:%M}. Internal mail for the Basis team only.</p>"
            "</body></html>")
    subject = f"[REPORT HELD] InfraBeatOps -- {name}: monitoring incomplete" + \
              (f", retry at {due:%H:%M}" if due else ", not sent")
    try:
        # The end-of-chain "NOT sent" notice has its own kind: "report_held" is
        # switched off by default (IBO_DISABLED_MAIL_KINDS) and the last word
        # on a lost sheet must always reach Basis (review 28.10.2026).
        _send_internal(system_config, subject, html, "report_held" if due else "report_not_sent",
                       f"{now:%Y%m%d%H%M}", "WARNING")
    except Exception as exc:  # noqa: BLE001
        log.error(f"{name}: report-held mail failed: {exc}")
    try:
        from core import timeline
        timeline.add(name, "sweep", "Monitoring sheet held back (run incomplete)",
                     "; ".join(gate["reasons"] + _detail_lines(gate))[:900]
                     + (f"; retry at {due:%H:%M}" if due else "; not retried"),
                     severity="warning")
    except Exception:  # noqa: BLE001
        pass
    log.warning(f"{name}: report NOT sent to the client: {'; '.join(gate['reasons'] + _detail_lines(gate))}"
                + (f" -- retry at {due:%H:%M}" if due else ""))


def sent(system_config: dict, cycle: str, client_to: list[str], now: datetime | None = None) -> None:
    """The report went to the client: clear retries, tell Basis with an Acknowledge button."""
    from html import escape as e
    name = system_config.get("name", "")
    now = now or datetime.now()
    clear(name)
    try:
        from core import timeline
        timeline.add(name, "sweep", "Monitoring sheet sent to the client",
                     f"to {', '.join(client_to)}", severity="ok")
    except Exception:  # noqa: BLE001
        pass
    button = ""
    try:
        from core import action_links
        url = action_links.public_url()
        if url and action_links.internal_only(_basis(system_config)):
            token = action_links.make("ack_report", name, kind="report", ttl_hours=72,
                                      extra={"c": str(cycle)[:40]})
            button = ("<p>" + action_links._btn(f"{url}/access/action?t={token}",
                                                "Acknowledge: report checked", "#2563eb")
                      + "<br><span style='color:#64748b;font-size:12px'>Records on the Timeline who "
                        "checked this report. Valid 72 hours, works once.</span></p>")
        else:
            button = ("<p style='color:#64748b'>No Acknowledge button: IBO_PUBLIC_URL and "
                      "IBO_INTERNAL_DOMAINS must be set in .env.</p>")
    except Exception as exc:  # noqa: BLE001
        log.info(f"{name}: acknowledge button skipped: {exc}")
    html = ("<html><body style='font-family:Segoe UI,Arial,sans-serif;color:#0f172a'>"
            f"<h2 style='color:#16a34a;margin:0 0 10px'>Monitoring sheet sent: {e(name)}</h2>"
            f"<p>The complete monitoring report (PDF and Excel) of {e(str(cycle))} was sent to "
            f"{e(', '.join(client_to) or 'the client')}.</p>{button}"
            "<p style='color:#64748b'>Internal mail for the Basis team only.</p></body></html>")
    try:
        _send_internal(system_config, f"[REPORT SENT] InfraBeatOps -- {name}: please check and acknowledge",
                       html, "report_sent", str(cycle), "INFO")
    except Exception as exc:  # noqa: BLE001
        log.error(f"{name}: report-sent notice failed: {exc}")


def acknowledge(system: str, cycle: str, by: str) -> str:
    from core import timeline
    timeline.add(system, "sweep", "Monitoring report checked", f"report of {cycle}; by {by}", severity="ok")
    return f"Done. The {system} monitoring report of {cycle} is recorded as checked by {by}."
