from pathlib import Path
import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
)
from core.models import MonitoringResult
from reporting.system_paths import system_report_dir, system_excel_path, system_pdf_path, system_evidence_root
from utils.paths import BASE_DIR
from utils.logger import get_logger

log = get_logger(__name__, "application")

def _text(v):
    return str(getattr(v, "value", v) if v is not None else "")

def _real_path(value):
    p = Path(value)
    return str(p if p.is_absolute() else Path(BASE_DIR) / p)

def _incidents(result):
    rows = []
    for item in result.incidents or []:
        rows.append(dict(item) if isinstance(item, dict) else getattr(item, "__dict__", {"incident": str(item)}))
    return rows

def generate_system_excel(result: MonitoringResult, gui_results=None, output_path=None):
    path = output_path or system_excel_path(result.system, result.cycle_timestamp)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "System Summary"
    summary = [
        ["InfraBeatOps Monitoring Report", ""],
        ["System", result.system],
        ["Client", result.client],
        ["Execution Time", result.cycle_timestamp.strftime("%Y-%m-%d %H:%M:%S")],
        ["Overall Status", _text(result.overall_status)],
        ["Metrics Evaluated", len(result.metrics)],
        ["Critical", len(result.critical_metrics())],
        ["Warning", len(result.warning_metrics())],
        ["Unknown", len(result.unknown_metrics())],
        ["T-Codes Executed", len(gui_results or [])],
        ["T-Codes Failed", sum(1 for m in (gui_results or []) if getattr(m, "display_value", "") == "failed")],
    ]
    for row in summary:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:B1")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60

    def sheet(name, headers, rows):
        sh = wb.create_sheet(name)
        sh.append(headers)
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="37474F")
        for row in rows:
            sh.append(row)
        sh.freeze_panes = "A2"
        for col in range(1, len(headers) + 1):
            vals = [len(str(sh.cell(r, col).value or "")) for r in range(1, sh.max_row + 1)]
            sh.column_dimensions[get_column_letter(col)].width = min(60, max(14, max(vals, default=14) + 2))
        return sh

    sheet("Metrics",
        ["Timestamp","System","Client","Metric","Value","Warning","Critical","Status","TCode","Detail","Screenshot"],
        [[result.cycle_timestamp.strftime("%Y-%m-%d %H:%M:%S"), result.system, result.client,
          m.name, m.value if m.value is not None else m.display_value, m.threshold_warning,
          m.threshold_critical, _text(m.status), m.tcode or "", m.detail, m.screenshot_path or ""]
         for m in result.metrics])

    sheet("TCode Results",
        ["Evidence ID","TCode","Status","Display Value","Attempts","Recovery","Detail","Screenshots"],
        [[(getattr(m,"extra_data",{}) or {}).get("evidence_id",""), m.tcode or "",
          _text(m.status), m.display_value, (getattr(m,"extra_data",{}) or {}).get("attempts",0),
          "; ".join((getattr(m,"extra_data",{}) or {}).get("recovery_actions",[]) or []),
          m.detail, "; ".join(m.screenshot_paths or [])] for m in (gui_results or [])])

    sheet("Evidence Index",
        ["Evidence ID","System","Client","TCode","Attempts","Status","Recovery","Screenshots"],
        [[(getattr(m,"extra_data",{}) or {}).get("evidence_id",""),
          (getattr(m,"extra_data",{}) or {}).get("system",result.system),
          (getattr(m,"extra_data",{}) or {}).get("client",result.client),
          m.tcode or "", (getattr(m,"extra_data",{}) or {}).get("attempts",0), _text(m.status),
          "; ".join((getattr(m,"extra_data",{}) or {}).get("recovery_actions",[]) or []),
          "; ".join(m.screenshot_paths or [])] for m in (gui_results or [])])

    incidents = _incidents(result)
    sheet("Incidents",
        ["Incident ID","Severity","Title","Description","Root Cause","Recommended Action"],
        [[i.get("incident_id", i.get("id","")), i.get("severity", i.get("status","")),
          i.get("title", i.get("name","")), i.get("description", i.get("detail","")),
          i.get("root_cause", i.get("likely_root_cause","")),
          i.get("recommended_action", i.get("recommendation",""))] for i in incidents])

    recovery = []
    for m in gui_results or []:
        d = getattr(m, "extra_data", {}) or {}
        if d.get("attempts", 1) > 1 or d.get("recovery_actions"):
            recovery.append([d.get("evidence_id",""), m.tcode or "", d.get("attempts",0),
                             "; ".join(d.get("recovery_actions",[]) or []), m.display_value, m.detail])
    sheet("Recovery History", ["Evidence ID","TCode","Attempts","Recovery Actions","Final Result","Detail"], recovery)

    ai = result.ai_analysis
    sheet("AI Analysis", ["Severity","Likely Root Cause","Evidence","Recommended Actions","Confidence"],
          [[ai.severity, ai.likely_root_cause, "; ".join(ai.evidence or []),
            "; ".join(ai.recommended_actions or []), ai.confidence]] if ai else [])
    wb.save(path)
    return str(path)

# ---------------------------------------------------------------------------
# PDF report
#
# Layout: cover summary (status, counts, key findings) -> check results ->
# analysis -> screen-evidence appendix, one screenshot per check, two per
# page, with a small caption instead of a bold "<TCode> — Screenshot N"
# heading. Status words and colours are the Excel sheet's own.
# ---------------------------------------------------------------------------
import re as _re
from datetime import datetime as _dt
from xml.sax.saxutils import escape as _esc

from reportlab.pdfgen import canvas as _rl_canvas

_NAVY = colors.HexColor("#1F3A5F")
_INK = colors.HexColor("#1F2933")
_MUTED = colors.HexColor("#6B7785")
_RULE = colors.HexColor("#D5DBE1")
_ZEBRA = colors.HexColor("#F6F8FA")

# status -> (text colour, fill)
_STATUS_COLOURS = {
    "OK": ("#1B5E20", "#E6F4EA"),
    "WARNING": ("#7A4F01", "#FFF1C2"),
    "CRITICAL": ("#9B1C1C", "#FDE2E1"),
    "UNKNOWN": ("#44505C", "#ECEFF3"),
    "FAILED": ("#9B1C1C", "#FDE2E1"),
    "CONFIRM": ("#1F3A5F", "#E3ECF7"),
    "INCIDENT": ("#9B1C1C", "#FDE2E1"),
}
_STATUS_LABEL = {"UNKNOWN": "NOT MEASURED", "CONFIRM": "CROSS-CHECK"}


def _t(value) -> str:
    """Escape text for a reportlab Paragraph ("<ok>" in SAP text broke it)."""
    return _esc(str(value if value is not None else ""))


def _primary_screenshot(paths):
    """
    The one screen per check that shows the result. Selection screens and
    popup captures are skipped: SM37's "records passed" capture was taken
    with the popup open, but the capture grabs the main window, so it was
    the job list a second time.
    """
    existing = [p for p in (paths or []) if p and os.path.exists(p)]
    if not existing:
        return None
    preferred = [p for p in existing
                 if not _re.search(r"_(selection|records_passed)_", Path(p).name, _re.I)]
    return (preferred or existing)[-1]


def _styles():
    base = getSampleStyleSheet()["BodyText"]
    s = {}
    s["title"] = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=19,
                                leading=23, textColor=_NAVY, spaceAfter=2)
    s["subtitle"] = ParagraphStyle("st", parent=base, fontSize=10, leading=13, textColor=_MUTED)
    s["h1"] = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=12.5,
                             leading=15, textColor=_NAVY, spaceBefore=12, spaceAfter=6)
    s["h2"] = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=9.5,
                             leading=12, textColor=_INK, spaceBefore=8, spaceAfter=4)
    s["body"] = ParagraphStyle("b", parent=base, fontSize=8.8, leading=12, textColor=_INK)
    s["muted"] = ParagraphStyle("m", parent=s["body"], fontSize=7.8, leading=10, textColor=_MUTED)
    s["cell"] = ParagraphStyle("c", parent=base, fontSize=7.6, leading=9.6, textColor=_INK)
    s["cellb"] = ParagraphStyle("cb", parent=s["cell"], fontName="Helvetica-Bold")
    s["head"] = ParagraphStyle("hd", parent=s["cell"], fontName="Helvetica-Bold",
                               textColor=colors.white)
    s["label"] = ParagraphStyle("l", parent=s["cell"], textColor=_MUTED)
    s["caption"] = ParagraphStyle("cap", parent=base, fontSize=7.4, leading=9.5,
                                  textColor=_MUTED, alignment=1)
    s["kpi"] = ParagraphStyle("k", parent=base, fontName="Helvetica-Bold", fontSize=17,
                              leading=20, alignment=1)
    s["kpil"] = ParagraphStyle("kl", parent=base, fontSize=7.2, leading=9, alignment=1,
                               textColor=_MUTED)
    s["item"] = ParagraphStyle("i", parent=s["body"], leftIndent=0, spaceAfter=3)
    return s


def _ai_unavailable(ai) -> bool:
    """The analyzer's placeholder when every AI provider failed (quota, bad keys)."""
    return (str(getattr(ai, "root_cause_category", "") or "").upper() == "AI_UNAVAILABLE"
            or str(getattr(ai, "likely_root_cause", "") or "").startswith(
                "The deterministic monitoring and correlation engines completed"))


def _code(tcode) -> str:
    """Display T-code: SM37_CANCELLED is a second SM37 check, not a T-code."""
    return str(tcode or "").split("_")[0] if str(tcode or "").startswith("SM37_") else str(tcode or "")


def _badge(status, styles):
    fg, _ = _STATUS_COLOURS.get(status, _STATUS_COLOURS["UNKNOWN"])
    label = _STATUS_LABEL.get(status, status)
    return Paragraph(f'<font color="{fg}"><b>{_t(label)}</b></font>', styles["cell"])


def _make_canvas(meta):
    class _ReportCanvas(_rl_canvas.Canvas):
        """Header band, footer and 'Page X of Y' on every page."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total):
            width, height = A4
            left, right = 1.6 * cm, width - 1.6 * cm
            self.saveState()
            self.setFont("Helvetica-Bold", 8.5)
            self.setFillColor(_NAVY)
            self.drawString(left, height - 1.15 * cm, "InfraBeatOps")
            self.setFont("Helvetica", 8.5)
            self.setFillColor(_MUTED)
            self.drawString(left + 2.05 * cm, height - 1.15 * cm, "SAP BASIS Monitoring Report")
            self.drawRightString(right, height - 1.15 * cm, meta["header_right"])
            self.setStrokeColor(_RULE)
            self.setLineWidth(0.6)
            self.line(left, height - 1.35 * cm, right, height - 1.35 * cm)
            self.line(left, 1.35 * cm, right, 1.35 * cm)
            self.setFont("Helvetica", 7.2)
            self.drawString(left, 0.95 * cm, meta["footer_left"])
            self.drawRightString(right, 0.95 * cm, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    return _ReportCanvas


def _grid_style(extra=()):
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _NAVY),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, _RULE),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        *extra,
    ])


def generate_system_pdf(result: MonitoringResult, gui_results=None, output_path=None):
    from reporting.check_narratives import narrate_all, deterministic_analysis
    from reporting.excel_template_writer import _system_config

    path = output_path or system_pdf_path(result.system, result.cycle_timestamp)
    Path(path).parent.mkdir(parents=True, exist_ok=True)

    st = _styles()
    narratives = narrate_all(gui_results)
    da = deterministic_analysis(result, narratives, gui_results)
    counts = da["counts"]
    status = da["severity"]
    ai = result.ai_analysis

    config = _system_config(result.system)
    sid = str(config.get("sap_system_id") or "")
    environment = str(config.get("environment") or "")
    run_at = result.cycle_timestamp
    now = _dt.now()

    meta = {
        "header_right": f"{result.system} · Client {result.client} · {run_at:%d.%m.%Y %H:%M}",
        "footer_left": f"Generated {now:%d.%m.%Y %H:%M} by InfraBeatOps · Confidential",
    }
    width = A4[0] - 3.2 * cm

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=2.0 * cm, bottomMargin=1.9 * cm,
                            title=f"{result.system} SAP BASIS Monitoring Report",
                            author="InfraBeatOps")
    story = []

    # ---- cover summary ------------------------------------------------------
    story.append(Paragraph("SAP BASIS Monitoring Report", st["title"]))
    story.append(Paragraph(f"System {_t(result.system)} · Client {_t(result.client)} · "
                           f"run of {run_at:%d.%m.%Y at %H:%M:%S}", st["subtitle"]))
    story.append(Spacer(1, 0.45 * cm))

    def lv(label, value):
        return [Paragraph(_t(label), st["label"]), Paragraph(_t(value or "—"), st["cellb"])]

    info = Table([
        lv("System", result.system) + lv("Client", result.client),
        lv("SID", sid) + lv("Environment", environment),
        lv("Run started", f"{run_at:%d.%m.%Y %H:%M:%S}") + lv("Report generated", f"{now:%d.%m.%Y %H:%M:%S}"),
        lv("Checks captured", f"{counts['captured']} of {counts['total']}")
        + lv("Evidence", "One screenshot per check (appendix)"),
    ], colWidths=[2.9 * cm, width / 2 - 2.9 * cm, 2.9 * cm, width / 2 - 2.9 * cm])
    info.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, _RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
    ]))
    story += [info, Spacer(1, 0.5 * cm)]

    fg, bg = _STATUS_COLOURS.get(status, _STATUS_COLOURS["UNKNOWN"])
    banner = Table([[
        Paragraph(f'<font color="{fg}" size="7.5">OVERALL STATUS</font><br/>'
                  f'<font color="{fg}" size="15"><b>{_t(_STATUS_LABEL.get(status, status))}</b></font>',
                  ParagraphStyle("bn", parent=st["body"], leading=17)),
        Paragraph(_t(da["headline"]), st["body"]),
    ]], colWidths=[4.4 * cm, width - 4.4 * cm])
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(bg)),
        ("LINEBEFORE", (0, 0), (0, 0), 4, colors.HexColor(fg)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    story += [banner, Spacer(1, 0.45 * cm)]

    tiles = [
        ("OK", counts["ok"], "Checks OK"),
        ("WARNING", counts["warning"], "Warnings"),
        ("CRITICAL", counts["critical"], "Critical"),
        ("UNKNOWN", counts["unknown"] + counts["failed"], "Not measured"),
        ("CONFIRM", counts["metric_findings"], f"RFC breaches ({counts['metric_critical']} critical)"),
    ]
    gap = 0.25 * cm
    tile_w = (width - gap * (len(tiles) - 1)) / len(tiles)
    row, widths, style = [], [], [("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                  ("TOPPADDING", (0, 0), (-1, -1), 7),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]
    for i, (key, number, label) in enumerate(tiles):
        tfg, tbg = _STATUS_COLOURS[key]
        row.append([Paragraph(f'<font color="{tfg}">{number}</font>', st["kpi"]),
                    Paragraph(_t(label), st["kpil"])])
        widths.append(tile_w)
        col = len(row) - 1
        style.append(("BACKGROUND", (col, 0), (col, 0), colors.HexColor(tbg)))
        if i < len(tiles) - 1:
            row.append("")
            widths.append(gap)
    kpi = Table([row], colWidths=widths)
    kpi.setStyle(TableStyle(style))
    story += [kpi]

    # Key findings: checks and metrics, worst first.
    rank = {"INCIDENT": -1, "CRITICAL": 0, "FAILED": 0, "WARNING": 1, "UNKNOWN": 2, "CONFIRM": 3}
    items = [(-1, "INCIDENT", _t(i)) for i in da.get("incidents", [])]
    checked = {n.tcode for n in narratives}
    for n in narratives:
        if n.status in ("CRITICAL", "FAILED", "WARNING"):
            items.append((rank[n.status], n.status,
                          f"<b>{_t(_code(n.tcode))}</b> {_t(n.task)} — {_t(n.result or n.observation)}"))
    for r in da["metric_rows"]:
        if r.get("source") == "screen" and r.get("tcode") in checked:
            continue  # the same reading as its check row
        items.append((rank[r["status"]], r["status"],
                      f"<b>{_t(r['name'])}</b> = {_t(r['value'])} ({_t(r['source'])})"))
    unknown = [n.task for n in narratives if n.status in ("UNKNOWN",)]
    if unknown:
        items.append((2, "UNKNOWN", "Not measured this run: <b>" + _t(", ".join(unknown))
                      + "</b> — see the screen evidence in the appendix."))
    for c in da["conflicts"]:
        items.append((3, "CONFIRM", _t(c)))
    items.sort(key=lambda x: x[0])

    story.append(Paragraph("Key findings", st["h1"]))
    if not items:
        story.append(Paragraph("No findings: every check read within range and no metric "
                               "breached a threshold.", st["body"]))
    else:
        shown = items[:9]
        rows = [[_badge(s_, st), Paragraph(text, st["cell"])] for _, s_, text in shown]
        t = Table(rows, colWidths=[2.3 * cm, width - 2.3 * cm])
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, _RULE),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(t)
        if len(items) > len(shown):
            story.append(Spacer(1, 0.15 * cm))
            story.append(Paragraph(f"{len(items) - len(shown)} further item(s) in the "
                                   "sections that follow.", st["muted"]))

    # ---- 1. check results -----------------------------------------------------
    story.append(PageBreak())
    story.append(Paragraph("1. Check results", st["h1"]))
    story.append(Paragraph(
        f"{counts['total']} SAP GUI checks run against {_t(result.system)} client "
        f"{_t(result.client)}. Status and figures match the Excel monitoring sheet "
        "row for row.", st["muted"]))
    story.append(Spacer(1, 0.2 * cm))

    head = [Paragraph(h, st["head"]) for h in ("T-code", "Check", "Status", "Result", "Observation")]
    rows, extra = [head], []
    for i, n in enumerate(narratives, start=1):
        result_text = n.result or n.value or ""
        observation = n.observation if n.observation and n.observation != result_text else ""
        rows.append([Paragraph(f"<b>{_t(_code(n.tcode))}</b>", st["cell"]), Paragraph(_t(n.task), st["cell"]),
                     _badge(n.status, st), Paragraph(_t(result_text), st["cell"]),
                     Paragraph(_t(observation), st["cell"])])
        _, sbg = _STATUS_COLOURS.get(n.status, _STATUS_COLOURS["UNKNOWN"])
        extra.append(("BACKGROUND", (2, i), (2, i), colors.HexColor(sbg)))
        if i % 2 == 0:
            extra.append(("BACKGROUND", (0, i), (1, i), _ZEBRA))
            extra.append(("BACKGROUND", (3, i), (4, i), _ZEBRA))
    table = Table(rows, colWidths=[1.6 * cm, 3.0 * cm, 2.2 * cm, 5.7 * cm, width - 12.5 * cm],
                  repeatRows=1)
    table.setStyle(_grid_style(extra))
    story.append(table)
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph("OK — within range · WARNING — review · CRITICAL — act now · "
                           "NOT MEASURED — the screen could not be read this run; see its "
                           "screenshot in the appendix.", st["muted"]))

    # ---- 2. analysis -----------------------------------------------------------
    story.append(Paragraph("2. Analysis", st["h1"]))
    story.append(Paragraph(f"<b>Status: {_t(_STATUS_LABEL.get(status, status))}.</b> "
                           f"{_t(da['headline'])}", st["body"]))

    if da.get("incidents"):
        story.append(Paragraph("Correlated incidents", st["h2"]))
        for i in da["incidents"]:
            story.append(Paragraph(f"• {_t(i)}", st["item"]))

    if da["metric_rows"]:
        story.append(Paragraph("Threshold breaches", st["h2"]))
        mrows = [[Paragraph(h, st["head"]) for h in ("Severity", "Metric", "Value", "Source", "Detail")]]
        mextra = []
        for i, r in enumerate(sorted(da["metric_rows"], key=lambda r: rank.get(r["status"], 9)), start=1):
            detail = r["detail"] if len(r["detail"]) <= 260 else r["detail"][:257] + "…"
            mrows.append([_badge(r["status"], st), Paragraph(_t(r["name"]), st["cell"]),
                          Paragraph(_t(r["value"]), st["cell"]), Paragraph(_t(r["source"]), st["cell"]),
                          Paragraph(_t(detail), st["cell"])])
            _, sbg = _STATUS_COLOURS[r["status"]]
            mextra.append(("BACKGROUND", (0, i), (0, i), colors.HexColor(sbg)))
        mt = Table(mrows, colWidths=[2.0 * cm, 4.3 * cm, 2.2 * cm, 1.7 * cm, width - 10.2 * cm],
                   repeatRows=1)
        mt.setStyle(_grid_style(mextra))
        story.append(mt)

    if da["conflicts"]:
        story.append(Paragraph("Cross-checks to confirm", st["h2"]))
        for c in da["conflicts"]:
            story.append(Paragraph(f"• {_t(c)}", st["item"]))

    if ai and _ai_unavailable(ai):
        # No model ran. The placeholder's "severity" and its "retry" advice
        # were being printed as if a model had rated the cycle.
        story.append(Paragraph("AI root-cause assessment", st["h2"]))
        story.append(Paragraph("AI analysis was unavailable this run; the findings above are "
                               "rules-based.", st["muted"]))
    elif ai:
        story.append(Paragraph("AI root-cause assessment", st["h2"]))
        conf = getattr(ai, "confidence", "") or ""
        model_sev = str(getattr(ai, "severity", "") or "")
        note = f"Model confidence: {_t(conf)}." if conf else ""
        if model_sev and model_sev.upper() != status:
            note += (f" The model rated this cycle {_t(model_sev.upper())}; the report status "
                     f"above comes from the measured thresholds and checks.")
        if getattr(ai, "likely_root_cause", ""):
            story.append(Paragraph(_t(ai.likely_root_cause), st["body"]))
        for a in (ai.recommended_actions or []):
            story.append(Paragraph(f"• {_t(a)}", st["item"]))
        if note:
            story.append(Paragraph(note.strip(), st["muted"]))

    if da["actions"]:
        story.append(Paragraph("Recommended actions", st["h2"]))
        for a in da["actions"]:
            story.append(Paragraph(f"• {_t(a)}", st["item"]))

    if not da["findings"]:
        story.append(Paragraph("No metric breached a threshold and every captured check read "
                               "normal.", st["body"]))

    # ---- appendix: screen evidence ---------------------------------------------
    figures = []
    for n in narratives:
        shot = _primary_screenshot(n.screenshots)
        if shot:
            figures.append((n, shot))

    story.append(PageBreak())
    story.append(Paragraph("Appendix — Screen evidence", st["h1"]))
    if not figures:
        story.append(Paragraph("No readable SAP GUI screenshot files were available for "
                               "embedding.", st["body"]))
    else:
        story.append(Paragraph("One screenshot per check, taken during the sweep. Full-resolution "
                               "files are retained on the collector under the Evidence ID in "
                               "each caption.", st["muted"]))
        story.append(Spacer(1, 0.3 * cm))
        for index, (n, shot) in enumerate(figures, start=1):
            try:
                from PIL import Image as PILImage
                with PILImage.open(shot) as im:
                    w_px, h_px = im.size
                scale = min(width / max(w_px, 1), 11.0 * cm / max(h_px, 1))
                img = Image(shot, width=w_px * scale, height=h_px * scale)
            except Exception:
                img = Image(shot, width=width, height=width * 0.53)
            framed = Table([[img]], colWidths=[img.drawWidth + 2])
            framed.setStyle(TableStyle([
                ("BOX", (0, 0), (-1, -1), 0.5, _RULE),
                ("LEFTPADDING", (0, 0), (-1, -1), 1), ("RIGHTPADDING", (0, 0), (-1, -1), 1),
                ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]))
            captured = ""
            stamp = str(getattr(n, "captured_at", "") or "")
            if len(stamp) >= 19:
                captured = f" · captured {stamp[11:19]}"
            caption = (f"Figure {index} · {_t(_code(n.tcode))} — {_t(n.task)}{captured}"
                       + (f" · {_t(n.evidence_id)}" if n.evidence_id else ""))
            story.append(KeepTogether([framed, Spacer(1, 0.12 * cm),
                                       Paragraph(caption, st["caption"]),
                                       Spacer(1, 0.55 * cm)]))

    doc.build(story, canvasmaker=_make_canvas(meta))
    return str(path)

def generate_system_reports(result: MonitoringResult, gui_results=None):
    return {
        "excel": generate_system_excel(result, gui_results),
        "pdf": generate_system_pdf(result, gui_results),
        "report_dir": str(system_report_dir(result.system, result.cycle_timestamp)),
        "evidence_root": str(system_evidence_root(result.system, result.cycle_timestamp)),
    }
